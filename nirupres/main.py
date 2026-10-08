"""NIRUPRES 2.0.0 — extracted application subsystem.

Refactored from the verified 1.7.3 baseline. Existing project and Markdown
format versions remain unchanged.
"""
import json, sys, zipfile, tempfile, shutil, uuid, re, stat, os, platform, hashlib
from collections import OrderedDict
from pathlib import Path
from PySide6.QtCore import Qt, QTimer, QSize, Signal, QRectF, QElapsedTimer, QEvent, QPropertyAnimation, QEasingCurve, QParallelAnimationGroup, QStandardPaths, QUrl, QPointF, QVariantAnimation
from PySide6.QtGui import QAction, QKeySequence, QShortcut, QPixmap, QPainter, QColor, QFont, QFontMetrics, QFontDatabase, QPdfWriter, QPageSize, QPageLayout, QImage, QImageReader, QIcon, QTextDocument, QAbstractTextDocumentLayout, QPainterPath
from PySide6.QtWidgets import (QApplication,QMainWindow,QWidget,QVBoxLayout,QHBoxLayout,QListWidget,QListWidgetItem,QLineEdit,QTextEdit,QComboBox,QPushButton,QLabel,QFileDialog,QSplitter,QMessageBox,QToolBar,QCheckBox,QDialog,QGridLayout,QDialogButtonBox,QGraphicsScene,QGraphicsPixmapItem,QGraphicsBlurEffect,QSlider,QGroupBox,QFormLayout,QMenu,QInputDialog,QGraphicsOpacityEffect)

from .foundation import *
from .foundation import _safe_zip_name, _validate_project_data, _cached_pixmap, _image_dimensions, _optimize_asset, _paired_font, _parse_two_column_body, _plain_markup, _inline_html, _format_markdown_selection, _IMAGE_CACHE, _IMAGE_CACHE_BYTES
from .presentation import *
from .presentation import _web_video_url
from .editor import *
from .markdown_io import *
from .markdown_io import _md_bool, _meta_block
from .markdown_io import _md_bool, _meta_block
class Main(QMainWindow):
    def __init__(self):
        super().__init__(); DATA.mkdir(parents=True,exist_ok=True); self.data=clone(DEFAULT); self.path=None; self.dirty=False; self.presenter=None; self.extract=tempfile.TemporaryDirectory(prefix='nirupres-'); self.history=[]; self.history_pos=-1; self._restoring=False; self._slide_clip=[]; self.outline_mode=False; self.file_prefs=load_file_prefs()
        self.setWindowTitle(f'{APP} {VERSION}'); self.resize(1540,920); self.build(); self.setStyleSheet(STYLE); self.reload_recents(); self.refresh(); self.push_history()
        self.autotimer=QTimer(self); self.autotimer.setSingleShot(True); self.autotimer.setInterval(2500); self.autotimer.timeout.connect(self.autosave); self.histtimer=QTimer(self); self.histtimer.setSingleShot(True); self.histtimer.setInterval(500); self.histtimer.timeout.connect(self.push_history); QTimer.singleShot(250,self.offer_recovery)
    def action(self,t,slot,key=None): a=QAction(t,self); a.triggered.connect(slot); key and a.setShortcut(QKeySequence(key)); return a
    def build(self):
        tb=QToolBar(); tb.setMovable(False); self.addToolBar(tb); self.toolbar=tb
        for a in [self.action('NEW',self.new,'Ctrl+N'),self.action('OPEN',self.open,'Ctrl+O'),self.action('SAVE',self.save,'Ctrl+S')]: tb.addAction(a)
        self.addAction(self.action('SAVE AS',self.save_as,'Ctrl+Shift+S'))
        tb.addSeparator(); tb.addAction(self.action('+ SLIDE',self.add,'Ctrl+Enter')); tb.addAction(self.action('DUPLICATE',self.dup,'Ctrl+D')); tb.addAction(self.action('DELETE',self.delete,'Delete')); tb.addSeparator(); tb.addAction(self.action('↶',self.undo,'Ctrl+Z'));tb.addAction(self.action('↷',self.redo,'Ctrl+Shift+Z'))
        tb.addSeparator(); export=QPushButton('EXPORT'); export.setObjectName('menuButton'); export.setToolTip('Export presentation…'); em=QMenu(export); em.addAction(self.action('Current slide → PNG',self.export_png,'Ctrl+Alt+E'));em.addAction(self.action('All slides → PNG',self.export_all_png));em.addAction(self.action('Deck → PDF',self.export_pdf,'Ctrl+Shift+E'));em.addAction(self.action('Deck → PPTX (pixel-perfect)',self.export_pptx)); em.addSeparator(); em.addAction(self.action('NIRUPRES Markdown…',self.export_markdown)); em.addAction(self.action('AI brief (current deck)…',self.export_ai_brief)); em.addAction(self.action('Copy AI prompt',self.copy_ai_prompt)); em.addAction(self.action('AI template/spec…',self.export_ai_template));export.setMenu(em);tb.addWidget(export); tb.addAction(self.action('IMPORT MD',self.import_markdown)); tb.addSeparator()
        tb.addAction(self.action('FOCUS',self.toggle_focus,'Tab')); tb.addAction(self.action('⌘',self.command_palette,'Ctrl+K')); about=QPushButton('?'); about.setFixedWidth(34); about.setToolTip('About NIRUPRES / diagnostics'); about.clicked.connect(self.about); tb.addWidget(about); tb.addSeparator()
        present=QPushButton('▶ PRESENT'); present.clicked.connect(lambda:self.present(False)); present.setObjectName('present'); present.setToolTip('Present from current slide · F5'); tb.addWidget(present)
        pstart=QPushButton('|▶'); pstart.setObjectName('presentAux'); pstart.setFixedWidth(42); pstart.setToolTip('Present from beginning · Shift+F5'); pstart.clicked.connect(lambda:self.present(True)); tb.addWidget(pstart)
        pview=QPushButton('▣'); pview.setObjectName('presentAux'); pview.setFixedWidth(42); pview.setToolTip('Presenter View / display'); pview.clicked.connect(self.presenter_view); tb.addWidget(pview)
        online=QPushButton('◉ ONLINE'); online.setObjectName('presentAux'); online.setMinimumWidth(86); online.setToolTip('Online presentation · opens a clean Audience window for Teams/Meet/Zoom sharing'); online.clicked.connect(lambda:self.online_present(False)); tb.addWidget(online)
        pover=QPushButton('▦'); pover.setObjectName('presentAux'); pover.setFixedWidth(42); pover.setToolTip('Presentation overview'); pover.clicked.connect(self.presentation_overview); tb.addWidget(pover)
        # Hidden but keyboard-accessible editing actions
        for a in [self.action('COPY SLIDE',self.copy_slides,'Ctrl+Shift+C'),self.action('PASTE SLIDE',self.paste_slides,'Ctrl+Shift+V'),self.action('SAME LAYOUT',self.add_same_layout,'Ctrl+Shift+D')]: self.addAction(a)
        self.addAction(self.action('PRESENT F5',lambda:self.present(False),'F5')); self.addAction(self.action('PRESENT FROM START',lambda:self.present(True),'Shift+F5'))
        root=QSplitter(); self.root_splitter=root; self.setCentralWidget(root)
        left=QWidget(); self.left_panel=left; lv=QVBoxLayout(left); lv.setContentsMargins(10,10,6,10); hh=QHBoxLayout(); hh.addWidget(QLabel('SLIDES')); hh.addStretch(); self.count=QLabel(); hh.addWidget(self.count); self.outline_btn=QPushButton('OUTLINE'); self.outline_btn.setMinimumWidth(72); self.outline_btn.setToolTip('Switch between slide thumbnails and compact outline'); self.outline_btn.setCheckable(True); self.outline_btn.toggled.connect(self.toggle_outline); hh.addWidget(self.outline_btn); collapse=QPushButton('‹');collapse.setFixedWidth(28);collapse.setToolTip('Hide Slides panel');collapse.clicked.connect(lambda:self.set_side_panel('left',False));hh.addWidget(collapse);lv.addLayout(hh); self.list=SlideList(); self.list.setSpacing(5); self.list.currentRowChanged.connect(self.select); self.list.reordered.connect(self.drag_reorder); lv.addWidget(self.list); plus=QPushButton('+ ADD SLIDE');plus.clicked.connect(self.add);lv.addWidget(plus); presets=QPushButton('PRESETS');presets.setToolTip('Insert a ready-made slide pattern');pmenu=QMenu(presets);
        for label,key in [('Uddevalla title','TITLE'),('Agenda','BULLETS'),('Section divider','SECTION'),('Key message','STATEMENT'),('Three points','BULLETS'),('Image + message','SPLIT'),('Result / KPI','BIG NUMBER'),('Thank you','END')]: pmenu.addAction(label,lambda _=False,k=key:self.add_preset(k));
        presets.setMenu(pmenu);lv.addWidget(presets); root.addWidget(left)
        center=QWidget(); cv=QVBoxLayout(center); cv.setContentsMargins(12,10,12,10)
        top=QHBoxLayout(); self.showleft=QPushButton('SLIDES');self.showleft.setMinimumWidth(70);self.showleft.setCheckable(True);self.showleft.setChecked(True);self.showleft.setToolTip('Show or hide the Slides panel');self.showleft.toggled.connect(lambda on:self.set_side_panel('left',on,from_toggle=True));top.addWidget(self.showleft); self.recent=QComboBox(); self.recent.setMinimumWidth(170); self.recent.setToolTip('Open a recent presentation'); self.recent.activated.connect(self.open_recent); top.addWidget(self.recent); top.addStretch(1); self.doc=QPushButton('Untitled  •');self.doc.setObjectName('doc');self.doc.clicked.connect(self.deck_settings);self.doc.setToolTip('Unsaved presentation · click for deck settings');top.addWidget(self.doc,0,Qt.AlignCenter); top.addStretch(1); self.showright=QPushButton('INSPECTOR');self.showright.setMinimumWidth(86);self.showright.setCheckable(True);self.showright.setChecked(True);self.showright.setToolTip('Show or hide the Inspector panel');self.showright.toggled.connect(lambda on:self.set_side_panel('right',on,from_toggle=True));top.addWidget(self.showright); self.state=QLabel(); self.state.setVisible(False); cv.addLayout(top)
        # Deck controls use semantic groups rather than a flat grid: a label is
        # always visually attached to the control it describes, while larger
        # gaps separate independent settings. This keeps the studio chrome calm
        # and prevents TRANSITION/NUMBERS or FONT/THEME from reading as pairs.
        deck=QVBoxLayout(); deck.setSpacing(7)
        deck_primary=QHBoxLayout(); deck_primary.setSpacing(8)
        deck_primary.addWidget(QLabel('THEME',objectName='deckLabel')); self.theme=QComboBox();self.theme.addItems(THEMES.keys());decorate_theme_combo(self.theme);self.theme.setMinimumWidth(170);self.theme.setToolTip('Deck theme');deck_primary.addWidget(self.theme,2)
        deck_primary.addSpacing(16); deck_primary.addWidget(QLabel('FONT',objectName='deckLabel'));self.deckfont=QComboBox();self.deckfont.addItems(QFontDatabase.families());self.deckfont.setMinimumWidth(220);self.deckfont.setToolTip('Deck typeface');deck_primary.addWidget(self.deckfont,3)
        deck_primary.addSpacing(16); deck_primary.addWidget(QLabel('RATIO',objectName='deckLabel'));self.aspect=QComboBox();self.aspect.addItems(['16:9','16:10','4:3','A4']);self.aspect.setMinimumWidth(96);self.aspect.setMaximumWidth(118);self.aspect.setToolTip('Presentation aspect ratio');deck_primary.addWidget(self.aspect); deck_primary.addStretch(1); deck.addLayout(deck_primary)
        deck_output=QHBoxLayout(); deck_output.setSpacing(10)
        deck_output.addWidget(QLabel('TRANSITION',objectName='deckLabel')); self.transition=QComboBox(); self.transition.addItems(TRANSITIONS); self.transition.setMinimumWidth(128); self.transition.setMaximumWidth(160); self.transition.setToolTip('Audience transition: NONE, FADE, MORPH or REVEAL'); deck_output.addWidget(self.transition)
        deck_output.addSpacing(18); self.numbers=QCheckBox('NUMBERS');self.numbers.setToolTip('Show slide numbers on audience output');deck_output.addWidget(self.numbers)
        self.guides=QCheckBox('GUIDES');self.guides.setToolTip('Show editor alignment guides · never shown to audience');deck_output.addWidget(self.guides)
        self.logo=QCheckBox('LOGO');self.logo.setToolTip('Master switch for the official Uddevalla kommun logo · slide AUTO/ON/OFF is in STYLE');deck_output.addWidget(self.logo); deck_output.addStretch(1); deck.addLayout(deck_output)
        deck_meta=QHBoxLayout(); deck_meta.setSpacing(8); deck_meta.addWidget(QLabel('FOOTER',objectName='deckLabel')); self.footer=QLineEdit(); self.footer.setPlaceholderText('Optional footer · e.g. Socialtjänsten · 2026-09-24'); self.footer.setToolTip('Optional deck footer · alignment and size are available in Deck Settings'); deck_meta.addWidget(self.footer,1)
        deck_meta.addSpacing(12); self.deck_settings_btn=QPushButton('DECK SETTINGS…'); self.deck_settings_btn.setToolTip('Open all deck-level settings in one dialog'); self.deck_settings_btn.clicked.connect(self.deck_settings); deck_meta.addWidget(self.deck_settings_btn); deck.addLayout(deck_meta); cv.addLayout(deck)
        self.view=SlideCanvas();self.view.editRequested.connect(self.edit_canvas_text);cv.addWidget(self.view,1); tip=QLabel('DOUBLE-CLICK TO EDIT  ·  CTRL+K COMMANDS  ·  F5 PRESENT  ·  ONLINE SHARES AUDIENCE');tip.setObjectName('hint');cv.addWidget(tip);root.addWidget(center)
        right=QWidget(); self.right_panel=right; rv=QVBoxLayout(right);rv.setContentsMargins(6,10,10,10);rh=QHBoxLayout();rh.addWidget(QLabel('INSPECTOR'));rh.addStretch();rc=QPushButton('›');rc.setFixedWidth(28);rc.setToolTip('Hide Inspector panel');rc.clicked.connect(lambda:self.set_side_panel('right',False));rh.addWidget(rc);rv.addLayout(rh)
        content=QGroupBox('▼  CONTENT');content.setCheckable(True);content.setChecked(True);cf=QFormLayout(content);self.layout=QComboBox();self.layout.addItems(LAYOUTS);cf.addRow('Layout',self.layout);self.title=QLineEdit();cf.addRow('Title',self.title);self.body=QTextEdit();self.body.setMinimumHeight(100);cf.addRow('Body',self.body);self.list_style=QComboBox();self.list_style.addItems(['NUMBERS','DOTS','DASHES','NONE']);self.list_style.setToolTip('Marker style for BULLETS and AGENDA layouts');cf.addRow('List style',self.list_style);rv.addWidget(content);content.toggled.connect(lambda on,g=content:self.collapse_group(g,on))
        image=QGroupBox('▼  IMAGE');self.image_group=image;image.setCheckable(True);image.setChecked(True);iv=QVBoxLayout(image);rr=QHBoxLayout();self.img=DropLine();self.img.setPlaceholderText('Drop image here or choose…');self.img.fileDropped.connect(self.set_image);b=QPushButton('CHOOSE');b.clicked.connect(self.choose_image);rr.addWidget(self.img,1);rr.addWidget(b);iv.addLayout(rr);ir=QHBoxLayout();self.imode=QComboBox();self.imode.addItems(['FIT','FILL']);self.imode.setToolTip('FIT shows the whole image · FILL crops to the image frame');self.treatment=QComboBox();self.treatment.addItems(['NATURAL','MONO','DIM','CONTRAST']);self.treatment.setToolTip('Non-destructive image treatment');self.mono=QCheckBox('MONO');self.mono.hide();clear=QPushButton('REMOVE');clear.clicked.connect(lambda:self.img.setText(''));ir.addWidget(QLabel('MODE',objectName='deckLabel'));ir.addWidget(self.imode);ir.addSpacing(8);ir.addWidget(QLabel('TREATMENT',objectName='deckLabel'));ir.addWidget(self.treatment,1);ir.addWidget(clear);iv.addLayout(ir);fr=QHBoxLayout();self.fx=QComboBox();self.fx.addItems(['LEFT','CENTER','RIGHT']);self.fy=QComboBox();self.fy.addItems(['TOP','CENTER','BOTTOM']);fr.addWidget(QLabel('FOCAL X',objectName='deckLabel'));fr.addWidget(self.fx);fr.addWidget(QLabel('FOCAL Y',objectName='deckLabel'));fr.addWidget(self.fy);iv.addLayout(fr); bglabel=QLabel('BACKGROUND IMAGE');bglabel.setObjectName('deckLabel');iv.addWidget(bglabel);bgr=QHBoxLayout();self.bgimg=DropLine();self.bgimg.setPlaceholderText('Optional full-slide background…');self.bgimg.fileDropped.connect(self.set_background_image);bgchoose=QPushButton('CHOOSE');bgchoose.clicked.connect(self.choose_background_image);bgclear=QPushButton('REMOVE');bgclear.clicked.connect(lambda:self.bgimg.setText(''));bgr.addWidget(self.bgimg,1);bgr.addWidget(bgchoose);bgr.addWidget(bgclear);iv.addLayout(bgr);bgopts=QHBoxLayout();self.bgmono=QCheckBox('MONO');self.bgdim=QSlider(Qt.Horizontal);self.bgdim.setRange(0,90);self.bgdim.setToolTip('Theme-coloured dim veil · 75% is a safe readable default');self.bgblur=QSlider(Qt.Horizontal);self.bgblur.setRange(0,20);self.bgblur.setToolTip('Background-only blur');bgopts.addWidget(self.bgmono);bgopts.addWidget(QLabel('DIM'));bgopts.addWidget(self.bgdim,1);bgopts.addWidget(QLabel('BLUR'));bgopts.addWidget(self.bgblur,1);iv.addLayout(bgopts);bgf=QHBoxLayout();self.bgfx=QComboBox();self.bgfx.addItems(['LEFT','CENTER','RIGHT']);self.bgfy=QComboBox();self.bgfy.addItems(['TOP','CENTER','BOTTOM']);bgf.addWidget(QLabel('BG FOCAL X',objectName='deckLabel'));bgf.addWidget(self.bgfx);bgf.addWidget(QLabel('Y',objectName='deckLabel'));bgf.addWidget(self.bgfy);iv.addLayout(bgf); bgmr=QHBoxLayout();bgmr.addWidget(QLabel('BG MOTION',objectName='deckLabel'));self.bgmotion=QComboBox();self.bgmotion.addItems(['NONE','KEN BURNS']);self.bgmotion.setToolTip('Subtle background-only pan during presentation · disabled by Reduce motion');bgmr.addWidget(self.bgmotion,1);iv.addLayout(bgmr);self.caption=QLineEdit();self.caption.setPlaceholderText('Caption / attribution');iv.addWidget(self.caption); self.image_request=QLineEdit(); self.image_request.setPlaceholderText('AI image request / desired image…'); iv.addWidget(self.image_request);rv.addWidget(image);image.toggled.connect(lambda on,g=image:self.collapse_group(g,on))
        video=QGroupBox('▶  VIDEO');self.video_group=video;video.setCheckable(True);video.setChecked(False);vv=QVBoxLayout(video);self.video_source=QLineEdit();self.video_source.setPlaceholderText('YouTube / Vimeo URL or local video path…');vv.addWidget(self.video_source);vpr=QHBoxLayout();self.video_poster=DropLine();self.video_poster.setPlaceholderText('Optional poster image…');vpchoose=QPushButton('POSTER…');vpchoose.clicked.connect(self.choose_video_poster);vlocal=QPushButton('LOCAL VIDEO…');vlocal.clicked.connect(self.choose_video);vpr.addWidget(self.video_poster,1);vpr.addWidget(vpchoose);vpr.addWidget(vlocal);vv.addLayout(vpr);vo=QHBoxLayout();self.video_autoplay=QCheckBox('AUTOPLAY');self.video_muted=QCheckBox('MUTED');vo.addWidget(self.video_autoplay);vo.addWidget(self.video_muted);vo.addStretch();vv.addLayout(vo);rv.addWidget(video);video.toggled.connect(lambda on,g=video:self.collapse_group(g,on));self.collapse_group(video,False)
        style=QGroupBox('▶  STYLE');self.style_group=style;style.setCheckable(True);style.setChecked(False);sf=QFormLayout(style); self.override_theme=QCheckBox('Override deck theme'); sf.addRow('',self.override_theme); self.slide_theme=QComboBox();self.slide_theme.addItems(list(THEMES.keys()));decorate_theme_combo(self.slide_theme); self.slide_theme.setEnabled(False); sf.addRow('Slide theme',self.slide_theme); self.logo_mode=QComboBox();self.logo_mode.addItems(['AUTO','ON','OFF']);self.logo_mode.setToolTip('AUTO follows layout rules · ON forces logo · OFF hides logo on this slide');sf.addRow('Logo',self.logo_mode); self.accent_override=QComboBox();self.accent_override.addItems(['AUTO','BLUE','GREEN','YELLOW','RED','PINK','PURPLE']);self.accent_override.setToolTip('Per-slide Uddevalla accent override');sf.addRow('Accent',self.accent_override); self.override_font=QCheckBox('Override deck font'); sf.addRow('',self.override_font); self.slide_font=QComboBox();self.slide_font.addItems(QFontDatabase.families()); self.slide_font.setEnabled(False); sf.addRow('Slide font',self.slide_font); self.override_theme.toggled.connect(self.slide_theme.setEnabled); self.override_font.toggled.connect(self.slide_font.setEnabled); tr=QHBoxLayout();self.fontscale=QComboBox();self.fontscale.addItems(['−','AUTO','+']);self.align=QComboBox();self.align.addItems(['LEFT','CENTER','RIGHT']);self.weight=QComboBox();self.weight.addItems(['AUTO','LIGHT','REGULAR','BOLD']);tr.addWidget(QLabel('Size'));tr.addWidget(self.fontscale);tr.addWidget(QLabel('Align'));tr.addWidget(self.align);tr.addWidget(QLabel('Weight'));tr.addWidget(self.weight);sf.addRow('Typography',tr)
        self.brightness=QSlider(Qt.Horizontal);self.brightness.setRange(-30,30);self.brightness.setValue(0);self.brightness.setToolTip('Image brightness · 0 = unchanged');sf.addRow('Brightness',self.brightness);self.contrast=QSlider(Qt.Horizontal);self.contrast.setRange(0,50);self.contrast.setToolTip('Image contrast · 0 = unchanged');sf.addRow('Contrast',self.contrast);self.overlay=QSlider(Qt.Horizontal);self.overlay.setRange(0,60);self.overlay.setToolTip('Dark overlay strength · 0 = none');sf.addRow('Overlay',self.overlay);self.blur=QSlider(Qt.Horizontal);self.blur.setRange(0,10);self.blur.setToolTip('Image blur · 0 = none');sf.addRow('Blur',self.blur); self.apply_style_btn=QPushButton('APPLY STYLE TO SELECTED'); self.apply_style_btn.setToolTip('Copy this slide’s visual style to all selected slides without replacing their content'); self.apply_style_btn.clicked.connect(self.apply_style_to_selected); sf.addRow('',self.apply_style_btn);rv.addWidget(style);style.toggled.connect(lambda on,g=style:self.collapse_group(g,on));self.collapse_group(style,False)
        pres=QGroupBox('▶  PRESENTATION');self.presentation_group=pres;pres.setCheckable(True);pres.setChecked(False);pv=QVBoxLayout(pres);self.hidden=QCheckBox('HIDDEN SLIDE — skip while presenting');pv.addWidget(self.hidden); transrow=QHBoxLayout(); transrow.addWidget(QLabel('TRANSITION',objectName='deckLabel')); self.slide_transition=QComboBox(); self.slide_transition.addItems(['DECK','NONE','FADE','MORPH','REVEAL']); self.slide_transition.setToolTip('Override the deck transition for this slide'); transrow.addWidget(self.slide_transition,1); self.transition_preview_btn=QPushButton('▶ PREVIEW'); self.transition_preview_btn.setToolTip('Preview the transition into the current slide'); self.transition_preview_btn.clicked.connect(self.preview_transition); transrow.addWidget(self.transition_preview_btn); pv.addLayout(transrow);pv.addWidget(QLabel('SPEAKER NOTES'));self.notes=QTextEdit();self.notes.setPlaceholderText('Private notes — never shown to audience');self.notes.setMinimumHeight(110);pv.addWidget(self.notes);rv.addWidget(pres);pres.toggled.connect(lambda on,g=pres:self.collapse_group(g,on));self.collapse_group(pres,False);rv.addStretch();root.addWidget(right);root.setSizes([210,1020,340])
        vals=(self.layout,self.title,self.body,self.list_style,self.img,self.imode,self.treatment,self.mono,self.fx,self.fy,self.bgimg,self.bgmono,self.bgdim,self.bgblur,self.bgfx,self.bgfy,self.bgmotion,self.video_source,self.video_poster,self.video_autoplay,self.video_muted,self.override_theme,self.slide_theme,self.logo_mode,self.accent_override,self.override_font,self.fontscale,self.align,self.weight,self.brightness,self.contrast,self.overlay,self.blur,self.slide_font,self.hidden,self.slide_transition,self.caption,self.image_request,self.notes)
        for w in vals:
            if isinstance(w,QComboBox):w.currentTextChanged.connect(self.changed)
            elif isinstance(w,QTextEdit):w.textChanged.connect(self.changed)
            elif isinstance(w,QCheckBox):w.toggled.connect(self.changed)
            elif isinstance(w,QSlider):w.valueChanged.connect(self.changed)
            else:w.textChanged.connect(self.changed)
        self.theme.currentTextChanged.connect(self.theme_changed); self.deckfont.currentTextChanged.connect(self.theme_changed); self.aspect.currentTextChanged.connect(self.theme_changed); self.numbers.toggled.connect(self.theme_changed); self.logo.toggled.connect(self.theme_changed); self.transition.currentTextChanged.connect(self.theme_changed); self.footer.textEdited.connect(self.footer_changed); self.guides.toggled.connect(self.guides_changed)
    def set_side_panel(self,which,visible,from_toggle=False):
        panel=self.left_panel if which=='left' else self.right_panel
        toggle=self.showleft if which=='left' else self.showright
        panel.setVisible(bool(visible))
        if not from_toggle:
            toggle.blockSignals(True); toggle.setChecked(bool(visible)); toggle.blockSignals(False)
        # A visible, named toggle always remains in the center header, so a hidden
        # side panel can never become unreachable.
        toggle.setToolTip(('Hide ' if visible else 'Show ')+('Slides panel' if which=='left' else 'Inspector panel'))

    def collapse_group(self,g,on):
        base=g.title().replace('▼  ','').replace('▶  ','')
        g.setTitle(('▼  ' if on else '▶  ')+base)
        for w in g.findChildren(QWidget, options=Qt.FindDirectChildrenOnly): w.setVisible(on)
        g.setMaximumHeight(16777215 if on else 34)
    def _dialog_start(self, filename='', prefer_current=True):
        folder=preferred_folder(self.file_prefs,self.path,prefer_current)
        return str(folder/filename) if filename else str(folder)
    def _remember_dialog(self,path,selected_filter=''):
        if path: remember_file_folder(self.file_prefs,path,selected_filter)
    def file_workflow_settings(self):
        d=FileWorkflowSettings(self,self.file_prefs)
        if d.exec()!=QDialog.Accepted:return
        self.file_prefs=d.values(); save_file_prefs(self.file_prefs); self.statusBar().showMessage('File workflow settings updated',1800)
    def deck_settings(self):
        d=DeckSettings(self,self.data)
        if d.exec()!=QDialog.Accepted:return
        self.data['theme']=d.theme.currentText();self.data['fontFamily']=d.font.currentText() or 'Noto Sans';self.data['aspect']=d.aspect.currentText();self.data['showNumbers']=d.numbers.isChecked();self.data['showLogo']=d.logo.isChecked();self.data['transition']=d.transition.currentText();self.data['reduceMotion']=d.reduce_motion.isChecked();self.data['cinematicOpen']=d.cinematic_open.isChecked();self.data['cinematicClose']=d.cinematic_close.isChecked();self.data['targetMinutes']=d.target_minutes.value();self.data['footerText']=d.footer.text().strip();self.data['footerAlign']=normalize_footer_align(d.footer_align.currentText());self.data['footerSize']=normalize_footer_size(d.footer_size.currentText());self.mark_dirty();self.refresh(self.list.currentRow());self.push_history();self.statusBar().showMessage('Deck settings updated',1800)
    def edit_canvas_text(self):
        s=self.current()
        if not s:return
        d=TextEditor(self,s.get('title',''),s.get('body',''),s.get('layout',''),s.get('leftColumn') if (s.get('sideStructured') or s.get('twoColumnStructured')) else None,s.get('rightColumn') if (s.get('sideStructured') or s.get('twoColumnStructured')) else None)
        if d.exec()!=QDialog.Accepted:return
        self.title.setText(d.title.text())
        if s.get('layout') in ('TWO COLUMN','COMPARE') and d.left_body is not None:
            left=d.left_body.toPlainText().strip(); right=d.right_body.toPlainText().strip()
            # Persist both structured data and a portable body mirror atomically.
            s['leftColumn']=left; s['rightColumn']=right; s['sideStructured']=True; s['twoColumnStructured']=(s.get('layout')=='TWO COLUMN')
            mirror=left + ('\n\n|||\n\n' if left or right else '') + right
            self.body.blockSignals(True); self.body.setPlainText(mirror); self.body.blockSignals(False); s['body']=mirror
            self.mark_dirty(); i=self.list.currentRow(); self.view.render(s,self.data['theme'],i,len(self.data['slides']),self.data.get('showNumbers',False),self.data.get('fontFamily','Noto Sans'),self.data.get('aspect','16:9'),self.data.get('showLogo',True),self.data.get('footerText',''),self.data.get('footerAlign','LEFT'),self.data.get('footerSize','MEDIUM')); self.update_thumb(i)
        else:
            self.body.setPlainText(d.body_text())
        self.title.setFocus();self.push_history()
    def command_palette(self):
        d=QDialog(self);d.setWindowTitle('NIRUPRES — Command Palette');d.resize(620,420);v=QVBoxLayout(d);q=QLineEdit();q.setPlaceholderText('Type a command…');v.addWidget(q);lst=QListWidget();v.addWidget(lst,1)
        cmds=[('Add slide',self.add),('Duplicate slide',self.dup),('Delete selected slides',self.delete),('Copy slide',self.copy_slides),('Paste slide',self.paste_slides),('Add same layout',self.add_same_layout),('Deck settings',self.deck_settings),('File workflow settings',self.file_workflow_settings),('Import NIRUPRES Markdown',self.import_markdown),('Export NIRUPRES Markdown',self.export_markdown),('Export AI brief',self.export_ai_brief),('Copy AI prompt',self.copy_ai_prompt),('Export AI template',self.export_ai_template),('Deck health',self.validate_deck),('Asset manager',self.asset_manager),('Presentation overview',self.presentation_overview),('Export current slide to PNG',self.export_png),('Export all slides to PNG',self.export_all_png),('Export deck to PDF',self.export_pdf),('Export pixel-perfect PPTX',self.export_pptx),('Focus mode',self.toggle_focus),('Present from current',lambda:self.present(False)),('Present from beginning',lambda:self.present(True)),('Online presentation',lambda:self.online_present(False)),('Online presentation from beginning',lambda:self.online_present(True)),('Presentation check',self.presentation_check),('Selected slides · Logo AUTO',lambda:self.apply_to_selected('logoMode','AUTO')),('Selected slides · Logo ON',lambda:self.apply_to_selected('logoMode','ON')),('Selected slides · Logo OFF',lambda:self.apply_to_selected('logoMode','OFF')),('Selected slides · Accent AUTO',lambda:self.apply_to_selected('accentOverride','AUTO')),('Selected slides · Accent BLUE',lambda:self.apply_to_selected('accentOverride','BLUE')),('Selected slides · Accent GREEN',lambda:self.apply_to_selected('accentOverride','GREEN')),('Selected slides · Accent YELLOW',lambda:self.apply_to_selected('accentOverride','YELLOW')),('Selected slides · Accent RED',lambda:self.apply_to_selected('accentOverride','RED')),('Selected slides · Accent PINK',lambda:self.apply_to_selected('accentOverride','PINK')),('Selected slides · Accent PURPLE',lambda:self.apply_to_selected('accentOverride','PURPLE')),('Save',self.save),('Save as',self.save_as),('Open',self.open),('About / diagnostics',self.about)]
        def fill():
            term=q.text().lower().strip();lst.clear()
            for name,fn in cmds:
                if not term or term in name.lower():it=QListWidgetItem(name);it.setData(Qt.UserRole,fn);lst.addItem(it)
            if lst.count():lst.setCurrentRow(0)
        def go(item=None):
            it=item or lst.currentItem()
            if it:fn=it.data(Qt.UserRole);d.accept();fn()
        q.textChanged.connect(fill);q.returnPressed.connect(go);lst.itemDoubleClicked.connect(go);fill();q.setFocus();d.exec()
    def current(self): i=self.list.currentRow(); return self.data['slides'][i] if 0<=i<len(self.data['slides']) else None
    def refresh(self,sel=0):
        self.list.blockSignals(True); self.list.clear(); total=len(self.data['slides']); th=self.data.get('theme','B&W')
        for i,s in enumerate(self.data['slides']):
            it=QListWidgetItem(); self.list.addItem(it)
            if self.outline_mode:
                flags=('  [HIDDEN]' if s.get('hidden') else '')+('  [IMAGE?]' if s.get('imageRequest') else '')
                it.setText(f'{i+1:02d}  {s.get("title") or "(untitled)"}\n     {s.get("layout","TITLE")}{flags}'); it.setSizeHint(QSize(196,52))
            else:
                it.setSizeHint(QSize(196,135)); self.list.setItemWidget(it,Thumb(s,th,i,total,self.data.get('showLogo',True),self.data.get('footerText',''),self.data.get('footerAlign','LEFT'),self.data.get('footerSize','MEDIUM')))
        self.list.blockSignals(False); self.count.setText(str(total)); self.theme.blockSignals(True);self.theme.setCurrentText(th);self.theme.blockSignals(False);self.numbers.blockSignals(True);self.numbers.setChecked(self.data.get('showNumbers',False));self.numbers.blockSignals(False);self.logo.blockSignals(True);self.logo.setChecked(self.data.get('showLogo',True));self.logo.setEnabled(THEMES.get(th,{}).get('uddevalla',False));self.logo.blockSignals(False); self.deckfont.blockSignals(True);self.deckfont.setCurrentText(self.data.get('fontFamily','Noto Sans'));self.deckfont.blockSignals(False);self.aspect.blockSignals(True);self.aspect.setCurrentText(self.data.get('aspect','16:9'));self.aspect.blockSignals(False); self.transition.blockSignals(True);self.transition.setCurrentText(normalize_transition(self.data.get('transition','NONE')));self.transition.blockSignals(False); self.footer.blockSignals(True);self.footer.setText(self.data.get('footerText',''));self.footer.blockSignals(False); self.list.setCurrentRow(max(0,min(sel,total-1))); self.list.scrollToItem(self.list.currentItem(),QListWidget.PositionAtCenter); self.select(self.list.currentRow()); self.update_title()
    def select(self,i):
        if i<0 or i>=len(self.data['slides']):return
        s=self.data['slides'][i]; vals=(self.layout,self.title,self.body,self.list_style,self.img,self.imode,self.treatment,self.mono,self.fx,self.fy,self.bgimg,self.bgmono,self.bgdim,self.bgblur,self.bgfx,self.bgfy,self.bgmotion,self.video_source,self.video_poster,self.video_autoplay,self.video_muted,self.override_theme,self.slide_theme,self.logo_mode,self.accent_override,self.override_font,self.fontscale,self.align,self.weight,self.brightness,self.contrast,self.overlay,self.blur,self.slide_font,self.hidden,self.slide_transition,self.caption,self.image_request,self.notes)
        for w in vals:w.blockSignals(True)
        self.layout.setCurrentText(s['layout']);self.title.setText(s['title']);self.body.setPlainText(s['body']);self.list_style.setCurrentText(s.get('listStyle','NUMBERS'));self.img.setText(s['image']);self.imode.setCurrentText(s['imageMode']);self.treatment.setCurrentText(str(s.get('imageTreatment','MONO' if s.get('mono') else 'NATURAL')).upper());self.mono.setChecked(self.treatment.currentText()=='MONO');self.fx.setCurrentText({0.0:'LEFT',.5:'CENTER',1.0:'RIGHT'}.get(s.get('focalX',.5),'CENTER'));self.fy.setCurrentText({0.0:'TOP',.5:'CENTER',1.0:'BOTTOM'}.get(s.get('focalY',.5),'CENTER'));self.bgimg.setText(s.get('backgroundImage',''));self.bgmono.setChecked(bool(s.get('backgroundMono',False)));self.bgdim.setValue(int(s.get('backgroundDim',75)));self.bgblur.setValue(int(s.get('backgroundBlur',0)));self.bgfx.setCurrentText({0.0:'LEFT',.5:'CENTER',1.0:'RIGHT'}.get(s.get('backgroundFocalX',.5),'CENTER'));self.bgfy.setCurrentText({0.0:'TOP',.5:'CENTER',1.0:'BOTTOM'}.get(s.get('backgroundFocalY',.5),'CENTER'));self.bgmotion.setCurrentText(s.get('backgroundMotion','NONE'));self.video_source.setText(s.get('videoSource',''));self.video_poster.setText(s.get('videoPoster',''));self.video_autoplay.setChecked(bool(s.get('videoAutoplay',False)));self.video_muted.setChecked(bool(s.get('videoMuted',False)));self.override_theme.setChecked(s.get('themeOverride','DECK')!='DECK');self.slide_theme.setEnabled(self.override_theme.isChecked());self.slide_theme.setCurrentText(s.get('themeOverride','DECK') if s.get('themeOverride','DECK')!='DECK' else self.data.get('theme','B&W'));self.logo_mode.setCurrentText(s.get('logoMode','AUTO'));self.accent_override.setCurrentText(s.get('accentOverride','AUTO'));self.override_font.setChecked(s.get('fontFamily','DECK')!='DECK');self.slide_font.setEnabled(self.override_font.isChecked());self.fontscale.setCurrentText({.85:'−',1.0:'AUTO',1.15:'+'}.get(float(s.get('fontScale',1.0)),'AUTO'));self.align.setCurrentText(s.get('align','LEFT'));self.weight.setCurrentText(s.get('weight','AUTO'));self.brightness.setValue(int(s.get('brightness',0)));self.contrast.setValue(int(s.get('contrast',0)));self.overlay.setValue(int(s.get('overlay',0)));self.blur.setValue(int(s.get('blur',0)));self.slide_font.setCurrentText(s.get('fontFamily','DECK') if s.get('fontFamily','DECK')!='DECK' else self.data.get('fontFamily','Noto Sans'));self.hidden.setChecked(bool(s.get('hidden',False)));self.slide_transition.setCurrentText(normalize_transition(s.get('transition','DECK'),True));self.caption.setText(s.get('caption',''));self.image_request.setText(s.get('imageRequest',''));self.notes.setPlainText(s['notes'])
        for w in vals:w.blockSignals(False)
        has_image=bool(s.get('image'))
        if has_image and not self.image_group.isChecked(): self.image_group.setChecked(True)
        self.view.render(s,self.data.get('theme','B&W'),i,len(self.data['slides']),self.data.get('showNumbers',False),self.data.get('fontFamily','Noto Sans'),self.data.get('aspect','16:9'),self.data.get('showLogo',True),self.data.get('footerText',''),self.data.get('footerAlign','LEFT'),self.data.get('footerSize','MEDIUM'))
    def changed(self,*_):
        if self._restoring:return
        s=self.current()
        if not s:return
        s.update(layout=self.layout.currentText(),title=self.title.text(),body=self.body.toPlainText(),listStyle=self.list_style.currentText(),image=self.img.text(),imageMode=self.imode.currentText(),imageTreatment=self.treatment.currentText(),mono=(self.treatment.currentText()=='MONO'),focalX={'LEFT':0.0,'CENTER':.5,'RIGHT':1.0}[self.fx.currentText()],focalY={'TOP':0.0,'CENTER':.5,'BOTTOM':1.0}[self.fy.currentText()],backgroundImage=self.bgimg.text(),backgroundMono=self.bgmono.isChecked(),backgroundDim=self.bgdim.value(),backgroundBlur=self.bgblur.value(),backgroundFocalX={'LEFT':0.0,'CENTER':.5,'RIGHT':1.0}[self.bgfx.currentText()],backgroundFocalY={'TOP':0.0,'CENTER':.5,'BOTTOM':1.0}[self.bgfy.currentText()],backgroundMotion=self.bgmotion.currentText(),videoSource=self.video_source.text(),videoPoster=self.video_poster.text(),videoAutoplay=self.video_autoplay.isChecked(),videoMuted=self.video_muted.isChecked(),themeOverride=(self.slide_theme.currentText() if self.override_theme.isChecked() else 'DECK'),logoMode=self.logo_mode.currentText(),accentOverride=self.accent_override.currentText(),fontScale={'−':.85,'AUTO':1.0,'+':1.15}[self.fontscale.currentText()],align=self.align.currentText(),weight=self.weight.currentText(),brightness=self.brightness.value(),contrast=self.contrast.value(),overlay=self.overlay.value(),blur=self.blur.value(),fontFamily=(self.slide_font.currentText() if self.override_font.isChecked() else 'DECK'),hidden=self.hidden.isChecked(),transition=self.slide_transition.currentText(),caption=self.caption.text(),imageRequest=self.image_request.text(),notes=self.notes.toPlainText()); self.mark_dirty(); i=self.list.currentRow(); self.view.render(s,self.data['theme'],i,len(self.data['slides']),self.data.get('showNumbers',False),self.data.get('fontFamily','Noto Sans'),self.data.get('aspect','16:9'),self.data.get('showLogo',True),self.data.get('footerText',''),self.data.get('footerAlign','LEFT'),self.data.get('footerSize','MEDIUM')); self.update_thumb(i); self.histtimer.start()
    def theme_changed(self,*_): self.data['theme']=self.theme.currentText();self.data['showNumbers']=self.numbers.isChecked();self.data['showLogo']=self.logo.isChecked();self.data['fontFamily']=self.deckfont.currentText() or 'Noto Sans';self.data['aspect']=self.aspect.currentText();self.data['transition']=self.transition.currentText();self.mark_dirty();self.refresh(self.list.currentRow());self.histtimer.start()
    def update_thumb(self,i):
        if i<0:return
        it=self.list.item(i)
        if not self.outline_mode:self.list.setItemWidget(it,Thumb(self.data['slides'][i],self.data['theme'],i,len(self.data['slides']),self.data.get('showLogo',True),self.data.get('footerText',''),self.data.get('footerAlign','LEFT'),self.data.get('footerSize','MEDIUM')))
        else:self.refresh(i)
    def footer_changed(self,text):
        if self._restoring:return
        self.data['footerText']=text.strip(); self.mark_dirty(); self.refresh_render_only(); self.histtimer.start()
    def refresh_render_only(self):
        i=self.list.currentRow(); s=self.current()
        if s:
            self.view.render(s,self.data.get('theme','B&W'),i,len(self.data['slides']),self.data.get('showNumbers',False),self.data.get('fontFamily','Noto Sans'),self.data.get('aspect','16:9'),self.data.get('showLogo',True),self.data.get('footerText',''),self.data.get('footerAlign','LEFT'),self.data.get('footerSize','MEDIUM'))
        for n in range(self.list.count()): self.update_thumb(n)
    def guides_changed(self,on): self.view.guides=on; self.view.update()
    def drag_reorder(self,old,new):
        if old<0 or new<0 or old==new:return
        slide=self.data['slides'].pop(old); self.data['slides'].insert(new,slide); self.refresh(new); self.mark_dirty(); self.push_history()
    def copy_slides(self):
        rows=sorted({x.row() for x in self.list.selectionModel().selectedRows()} or {self.list.currentRow()}); self._slide_clip=[clone(self.data['slides'][i]) for i in rows if i>=0]; self.statusBar().showMessage(f'Copied {len(self._slide_clip)} slide(s)',1800)
    def paste_slides(self):
        clips=getattr(self,'_slide_clip',[])
        if not clips:return
        at=self.list.currentRow()+1
        for n,x in enumerate(clips): x=slide_defaults(x); x['id']=str(uuid.uuid4()); self.data['slides'].insert(at+n,x)
        self.refresh(at); self.mark_dirty(); self.push_history()
    def save_as(self):
        old=self.path
        p,_=QFileDialog.getSaveFileName(self,'Save NIRUPRES As',self._dialog_start(self.path.name if self.path else 'presentation.nirupres'),'NIRUPRES (*.nirupres)')
        if not p:return False
        self._remember_dialog(p)
        self.path=Path(p if p.lower().endswith('.nirupres') else p+'.nirupres')
        if self.save():return True
        self.path=old; self.update_title(); return False
    def about(self):
        qtver=getattr(sys.modules.get('PySide6'),'__version__','unknown')
        screens='\n'.join(f'  {i+1}. {x.name()} — {x.geometry().width()}×{x.geometry().height()} @ {x.geometry().x()},{x.geometry().y()}' for i,x in enumerate(QApplication.screens())) or '  none'
        text=(f'NIRUPRES {VERSION}\n{TAGLINE}\n\n'
              f'Project format: {FORMAT} · NIRUPRES Markdown spec: {MD_FORMAT}\n'
              f'Python: {platform.python_version()}\nPySide6: {qtver}\n'
              f'Platform: {platform.system()} {platform.release()}\n'
              f'Session: {os.environ.get("XDG_SESSION_TYPE","unknown")} · Wayland: {os.environ.get("WAYLAND_DISPLAY","-")}\n'
              f'Desktop: {os.environ.get("XDG_CURRENT_DESKTOP","unknown")} · Portal: {os.environ.get("XDG_DESKTOP_PORTAL_DIR","auto")}\n'
              f'Screens:\n{screens}')
        d=QDialog(self); d.setWindowTitle(f'About {APP}'); d.resize(650,470); v=QVBoxLayout(d)
        title=QLabel('NIRUPRES'); title.setObjectName('aboutTitle'); v.addWidget(title); tag=QLabel(TAGLINE); tag.setWordWrap(True); tag.setObjectName('aboutTag'); v.addWidget(tag)
        box=QTextEdit(); box.setReadOnly(True); box.setPlainText(text); v.addWidget(box,1)
        row=QHBoxLayout(); copy=QPushButton('COPY DIAGNOSTICS'); close=QPushButton('CLOSE'); copy.clicked.connect(lambda:QApplication.clipboard().setText(text)); close.clicked.connect(d.accept); row.addWidget(copy); row.addStretch(); row.addWidget(close); v.addLayout(row); d.exec()
    def apply_style_to_selected(self):
        src=self.current()
        if not src:return
        rows=sorted({x.row() for x in self.list.selectionModel().selectedRows()})
        if len(rows)<2:
            self.statusBar().showMessage('Select two or more slides to apply style',2500); return
        keys=('themeOverride','logoMode','accentOverride','fontScale','align','weight','fontFamily','brightness','contrast','overlay','blur','imageMode','imageTreatment','mono','focalX','focalY','listStyle','backgroundImage','backgroundMono','backgroundDim','backgroundBlur','backgroundFocalX','backgroundFocalY','transition')
        for i in rows:
            if 0<=i<len(self.data['slides']):
                dst=self.data['slides'][i]
                for key in keys: dst[key]=clone(src.get(key))
        self.mark_dirty(); self.refresh(self.list.currentRow()); self.push_history(); self.statusBar().showMessage(f'Applied style to {len(rows)} selected slides',2500)

    def preview_transition(self):
        if not self.data.get('slides'):return
        cur=max(0,self.list.currentRow()); prev=cur-1
        while prev>=0 and self.data['slides'][prev].get('hidden',False): prev-=1
        if prev<0:
            self.statusBar().showMessage('Transition preview needs a previous visible slide',2500); return
        old=getattr(self,'_transition_preview',None)
        if old is not None and old.isVisible(): old.close()
        preview=Audience(self.data,prev,online=True); preview.setWindowTitle('NIRUPRES — Transition Preview'); preview.resize(960,540); preview.setAttribute(Qt.WA_DeleteOnClose,True); preview.show(); preview.raise_(); preview.activateWindow()
        self._transition_preview=preview
        QTimer.singleShot(250,lambda p=preview,n=cur: p.set_index(n) if p.isVisible() else None)

    def presentation_check(self):
        issues=self.validate_deck(silent=True) or []
        blocking=[x for x in issues if x[1] in ('ERROR','TODO')]
        if not issues:
            message_info(self,'Presentation Check','Ready to present.\n\nNo deck-health issues were found.'); return True
        msg=f'{len(issues)} note(s): {sum(x[1]=="ERROR" for x in issues)} error(s), {sum(x[1]=="WARN" for x in issues)} warning(s), {sum(x[1]=="TODO" for x in issues)} unresolved image request(s).'
        if blocking: msg+='\n\nThe presentation can still run, but audience output may be incomplete.'
        r=message_question(self,'Presentation Check',msg+'\n\nOpen Deck Health?',QMessageBox.Yes|QMessageBox.No,QMessageBox.No)
        if r==QMessageBox.Yes:self.validate_deck(False)
        return not blocking
    def export_png(self):
        s=self.current()
        if not s:return
        pth,_=QFileDialog.getSaveFileName(self,'Export current slide PNG',self._dialog_start(f'slide-{self.list.currentRow()+1:02d}.png'),'PNG (*.png)')
        if not pth:return
        self._remember_dialog(pth)
        if not pth.endswith('.png'):pth+='.png'
        img=self._render_export_image(s,self.list.currentRow(),*EXPORT_PROFILES['preview'])
        img.save(pth);self.statusBar().showMessage(f'PNG exported: {pth}',2500)
    def export_all_png(self):
        folder=QFileDialog.getExistingDirectory(self,'Export all slides as PNG',self._dialog_start())
        if not folder:return
        self._remember_dialog(folder)
        for i,s in enumerate(self.data['slides']):
            img=self._render_export_image(s,i,*EXPORT_PROFILES['preview'])
            img.save(str(Path(folder)/f'slide-{i+1:02d}.png'))
        self.statusBar().showMessage(f'Exported {len(self.data["slides"])} PNG slides',2500)
    def toggle_focus(self):
        entering=self.left_panel.isVisible() or self.right_panel.isVisible()
        self.set_side_panel('left',not entering); self.set_side_panel('right',not entering)
        for tb in self.findChildren(QToolBar):tb.setVisible(not entering)
    def reload_recents(self):
        if not hasattr(self,'recent'):return
        self.recent.blockSignals(True); self.recent.clear(); self.recent.addItem('RECENT FILES…','')
        try:
            for p in json.loads(RECENTS.read_text()) if RECENTS.exists() else []:
                if Path(p).exists():self.recent.addItem(Path(p).stem,p)
        except:pass
        self.recent.setCurrentIndex(0); self.recent.blockSignals(False)
    def open_recent(self,idx):
        p=self.recent.itemData(idx)
        if p:self.open_path(p)
        self.recent.setCurrentIndex(0)
    def open_path(self,p):
        if not self.confirm_discard():return
        candidate=tempfile.TemporaryDirectory(prefix='nirupres-')
        try:
            new_data=load_project(p,candidate.name)
            old_extract=self.extract; self.extract=candidate; self.data=new_data; self.path=Path(p); self.dirty=False; self.state.setText(''); self.history=[]; self.history_pos=-1
            self.refresh(); self.push_history(); self.clear_autosave(); self.add_recent(p); self._remember_dialog(p); old_extract.cleanup()
        except Exception as e:
            candidate.cleanup(); message_critical(self,'Open failed',f'NIRUPRES could not open this presentation.\n\n{e}')
    def add_recent(self,path):
        try:
            items=json.loads(RECENTS.read_text()) if RECENTS.exists() else []; p=str(Path(path)); items=[p]+[x for x in items if x!=p and Path(x).exists()]; RECENTS.write_text(json.dumps(items[:8]),encoding='utf-8'); self.reload_recents()
        except:pass
    def mark_dirty(self):
        self.dirty=True; self.state.setText('● UNSAVED'); self.update_title();
        if hasattr(self,'autotimer'): self.autotimer.start()
    def update_title(self):
        if self.path:
            display=self.path.name; tip=str(self.path)
        else:
            display=self.data.get('title','Untitled') or 'Untitled'; tip='Unsaved presentation'
        marker='  •' if self.dirty else ''
        fm=QFontMetrics(self.doc.font()); self.doc.setText(fm.elidedText(display,Qt.ElideMiddle,360)+marker)
        self.doc.setToolTip(tip+(' · unsaved changes' if self.dirty else '')+' · click for deck settings')
        self.setWindowTitle(f'{display}{" *" if self.dirty else ""} — {APP} {VERSION}')
    def push_history(self):
        if self._restoring:return
        snap=json.dumps(self.data,sort_keys=True,ensure_ascii=False)
        if self.history_pos>=0 and self.history[self.history_pos]==snap:return
        self.history=self.history[:self.history_pos+1]; self.history.append(snap); self.history=self.history[-60:]
        while len(self.history)>1 and sum(len(x.encode('utf-8')) for x in self.history)>MAX_HISTORY_BYTES: self.history.pop(0)
        self.history_pos=len(self.history)-1
    def restore_history(self,pos):
        if not (0<=pos<len(self.history)):return
        self._restoring=True;sel=self.list.currentRow();self.data=json.loads(self.history[pos]);self.history_pos=pos;self.refresh(sel);self._restoring=False;self.mark_dirty()
    def undo(self):self.restore_history(self.history_pos-1)
    def redo(self):self.restore_history(self.history_pos+1)
    def keyPressEvent(self,e):
        if e.modifiers() & Qt.ControlModifier and e.key()==Qt.Key_Up:self.move_slide(-1);return
        if e.modifiers() & Qt.ControlModifier and e.key()==Qt.Key_Down:self.move_slide(1);return
        super().keyPressEvent(e)
    def move_slide(self,d):
        i=self.list.currentRow();j=i+d
        if i<0 or j<0 or j>=len(self.data['slides']):return
        self.data['slides'][i],self.data['slides'][j]=self.data['slides'][j],self.data['slides'][i];self.refresh(j);self.mark_dirty();self.push_history()
    def add(self):
        d=LayoutPicker(self)
        if d.exec()!=QDialog.Accepted or not d.choice:return
        i=self.list.currentRow()+1; defaults={'TITLE':('Title','Subtitle'),'STATEMENT':('New statement','Short supporting line'),'TEXT':('Heading','Concise narrative text'),'BULLETS':('Key points','First point\nSecond point\nThird point'),'AGENDA':('Agenda','Opening\nContext\nDecision\nNext steps'),'IMAGE':('',''),'PHOTO':('Photo title','Short caption'),'HERO IMAGE':('A visual story','Short supporting line'),'SPLIT':('Heading','Supporting text'),'TWO COLUMN':('Heading','Left column\n\n|||\n\nRight column'),'COMPARE':('Compare','Current state\nKey limitation\n\n|||\n\nDesired state\nKey benefit'),'IMAGE + QUOTE':('A strong quote','Attribution'),'QUOTE':('Quote','Attribution'),'BIG NUMBER':('80%','Key result or context'),'DATA / KPI':('Results','80% | Faster\n12 | Systems\n47 | Risks\n2027 | Production'),'NUMBER GRID':('Key figures','80% | Faster\n12 | Systems\n47 | Risks\n2027 | Production'),'PROCESS':('Process','ANALYZE | Understand the need\nPILOT | Test safely\nDELIVER | Put into use\nFOLLOW UP | Measure and improve'),'TIMELINE':('Timeline','SEP | Pilot\nOCT | Test\nNOV | Release\nJAN | Production'),'MATRIX':('Framework','PEOPLE | Roles and capability\nPROCESS | Clear ways of working\nTECHNOLOGY | Fit-for-purpose tools\nGOVERNANCE | Ownership and control'),'TABLE':('Overview','AREA | STATUS | OWNER\nSecurity | Ready | IT\nTraining | In progress | HR\nLaunch | Planned | Team'),'SECTION':('Section','Short context'),'FULL BLEED':('Title','Subtitle'),'VIDEO':('Video','Optional caption'),'VIDEO + TEXT':('Video + context','Key context or talking points'),'END':('Thank you','Questions / contact')} ; t,b=defaults.get(d.choice,('',''));self.data['slides'].insert(i,slide_defaults({'layout':d.choice,'title':t,'body':b}));self.refresh(i);self.mark_dirty();self.push_history()
    def dup(self):
        s=self.current();i=self.list.currentRow()
        if s:self.data['slides'].insert(i+1,slide_defaults(clone(s)));self.refresh(i+1);self.mark_dirty();self.push_history()
    def delete(self):
        rows=sorted({x.row() for x in self.list.selectionModel().selectedRows()} or {self.list.currentRow()},reverse=True)
        if len(self.data['slides'])-len(rows)<1:return
        target=min(rows)
        for i in rows:
            if i>=0:self.data['slides'].pop(i)
        self.refresh(min(target,len(self.data['slides'])-1));self.mark_dirty();self.push_history()
    def add_same_layout(self):
        cur=self.current(); i=self.list.currentRow()+1; self.data['slides'].insert(i,slide_defaults({'layout':cur.get('layout','STATEMENT') if cur else 'STATEMENT','themeOverride':cur.get('themeOverride','DECK') if cur else 'DECK'}));self.refresh(i);self.mark_dirty();self.push_history()
    def add_preset(self,kind):
        presets={
            'TITLE':('TITLE','Presentation title','Short subtitle'),
            'BULLETS':('BULLETS','Agenda','First point\nSecond point\nThird point'),
            'SECTION':('SECTION','01 · Section','Short context'),
            'STATEMENT':('STATEMENT','One clear message','Short supporting sentence'),
            'SPLIT':('SPLIT','Image + message','Supporting text'),
            'BIG NUMBER':('BIG NUMBER','80%','Key result or KPI'),
            'VIDEO':('VIDEO','Video','Optional caption'),'VIDEO + TEXT':('VIDEO + TEXT','Video + context','Key context or talking points'),'END':('END','Thank you','Questions / contact')}; layout,title,body=presets.get(kind,presets['STATEMENT']); i=self.list.currentRow()+1; self.data['slides'].insert(i,slide_defaults({'layout':layout,'title':title,'body':body})); self.refresh(i);self.mark_dirty();self.push_history()
    def apply_to_selected(self,field,value):
        rows=sorted({x.row() for x in self.list.selectionModel().selectedRows()} or {self.list.currentRow()}); changed=0
        for i in rows:
            if 0<=i<len(self.data['slides']): self.data['slides'][i][field]=value; changed+=1
        if changed:self.refresh(rows[0]);self.mark_dirty();self.push_history();self.statusBar().showMessage(f'Applied to {changed} selected slide(s)',1800)
    def choose_image(self):
        p,_=QFileDialog.getOpenFileName(self,'Choose image',self._dialog_start(), 'Images (*.png *.jpg *.jpeg *.webp *.bmp)')
        if p:self._remember_dialog(p); self.set_image(p)
    def set_image(self,p):self.img.setText(p)
    def choose_background_image(self):
        p,_=QFileDialog.getOpenFileName(self,'Choose background image',self._dialog_start(), 'Images (*.png *.jpg *.jpeg *.webp *.bmp)')
        if p:self._remember_dialog(p); self.set_background_image(p)
    def set_background_image(self,p):self.bgimg.setText(p)
    def choose_video_poster(self):
        p,_=QFileDialog.getOpenFileName(self,'Choose video poster',self._dialog_start(),'Images (*.png *.jpg *.jpeg *.webp *.bmp)')
        if p:self._remember_dialog(p);self.video_poster.setText(p)
    def choose_video(self):
        p,_=QFileDialog.getOpenFileName(self,'Choose local video',self._dialog_start(),'Video (*.mp4 *.webm *.mov *.mkv)')
        if p:self._remember_dialog(p);self.video_source.setText(p)
    def new(self):
        if not self.confirm_discard():return
        d=QDialog(self); d.setWindowTitle('New presentation'); d.setMinimumWidth(520); v=QVBoxLayout(d); v.addWidget(QLabel('NEW PRESENTATION'))
        choice={'theme':'B&W','do_import':False}
        for label,theme in [('Blank','B&W'),('Uddevalla Blue','UDDEVALLA BLUE'),('Uddevalla Dark','UDDEVALLA DARK'),('Uddevalla Black','UDDEVALLA BLACK')]:
            b=QPushButton(label); b.clicked.connect(lambda _=False,t=theme:(choice.update(theme=t),d.accept())); v.addWidget(b)
        imp=QPushButton('FROM MARKDOWN / AI…'); imp.clicked.connect(lambda:(choice.update(do_import=True),d.accept())); v.addWidget(imp)
        cancel=QPushButton('CANCEL'); cancel.clicked.connect(d.reject); v.addWidget(cancel)
        if d.exec()!=QDialog.Accepted:return
        if choice['do_import']: self.import_markdown(); return
        self.data=clone(DEFAULT); self.data['theme']=choice['theme']; self.path=None;self.dirty=False;self.state.setText('');self.history=[];self.history_pos=-1;self.refresh();self.push_history()
    def open(self):
        p,_=QFileDialog.getOpenFileName(self,'Open NIRUPRES',self._dialog_start(prefer_current=False),'NIRUPRES (*.nirupres *.json)')
        if p:self._remember_dialog(p); self.open_path(p)
    def save(self):
        if not self.path:
            p,_=QFileDialog.getSaveFileName(self,'Save NIRUPRES',self._dialog_start('presentation.nirupres'),'NIRUPRES (*.nirupres)')
            if not p:return False
            self._remember_dialog(p)
            self.path=Path(p if p.lower().endswith('.nirupres') else p+'.nirupres')
        payload=clone(self.data);payload['format']=FORMAT;payload['appVersion']=VERSION
        tmp=None
        try:
            self.path.parent.mkdir(parents=True,exist_ok=True)
            used={}
            with tempfile.TemporaryDirectory(prefix='nirupres-save-') as td:
                ad=Path(td)/'assets';ad.mkdir()
                quality=self.file_prefs.get('assetQuality','OPTIMIZED'); asset_stats={'before':0,'after':0,'count':0}
                for s in payload['slides']:
                    for field in ('image','backgroundImage','videoPoster'):
                        src=Path(s.get(field,'')) if s.get(field) else None
                        if src and src.exists():
                            key=str(src.resolve())+'|'+quality
                            if key not in used:
                                dst,before_bytes,after_bytes,_,_=_optimize_asset(src,ad,quality)
                                used[key]=dst.name; asset_stats['before']+=before_bytes; asset_stats['after']+=after_bytes; asset_stats['count']+=1
                            s[field]='asset:'+used[key]
                for s in payload['slides']:
                    vsrc=str(s.get('videoSource','') or '')
                    if vsrc and not re.match(r'^[a-zA-Z]+://',vsrc):
                        src=Path(vsrc)
                        if src.exists() and src.is_file():
                            key='video|'+str(src.resolve())
                            if key not in used:
                                name=f'{hashlib.sha256(str(src.resolve()).encode()).hexdigest()[:10]}-{src.name}'; shutil.copy2(src,ad/name); used[key]=name
                            s['videoSource']='asset:'+used[key]
                project=Path(td)/'project.json'; project.write_text(json.dumps(payload,indent=2,ensure_ascii=False),encoding='utf-8')
                tmp=self.path.with_name(self.path.name+'.tmp')
                with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as z:
                    z.write(project,'project.json')
                    for f in ad.iterdir():z.write(f,'assets/'+f.name)
                # Verify the complete package before replacing the user's last good file.
                with zipfile.ZipFile(tmp) as z:
                    if z.testzip() is not None: raise ValueError('ZIP verification failed.')
                    probe=json.loads(z.read('project.json').decode('utf-8')); _validate_project_data(probe)
                os.replace(tmp,self.path)
            self.add_recent(self.path);self.dirty=False;self.state.setText('SAVED');self.update_title();self.clear_autosave();saved=max(0,asset_stats['before']-asset_stats['after']); self.statusBar().showMessage(f'Saved portable project · {asset_stats["count"]} image(s) · saved {saved/MB:.1f} MB: {self.path}',3200);return True
        except Exception as e:
            try:
                if tmp and Path(tmp).exists():Path(tmp).unlink()
            except Exception:pass
            message_critical(self,'Save failed',f'NIRUPRES could not save this presentation.\n\nYour previous file has not been replaced.\n\n{e}')
            self.mark_dirty(); return False
    def _render_export_image(self,s,i,width=1920,height=1080):
        ar={'16:9':16/9,'16:10':16/10,'4:3':4/3,'A4':297/210}.get(self.data.get('aspect','16:9'),16/9)
        if width/height>ar: width=int(height*ar)
        else: height=int(width/ar)
        img=QImage(width,height,QImage.Format.Format_ARGB32); img.fill(QColor(THEMES.get(self.data.get('theme','B&W'),THEMES['B&W'])['bg'])); q=QPainter(img)
        try: render_slide(q,QRectF(0,0,width,height),s,self.data.get('theme','B&W'),i,len(self.data['slides']),self.data.get('showNumbers',False),self.data.get('fontFamily','Noto Sans'),self.data.get('showLogo',True),self.data.get('footerText',''),self.data.get('footerAlign','LEFT'),self.data.get('footerSize','MEDIUM'))
        finally:
            if q.isActive(): q.end()
        return img
    def export_pptx(self):
        try:
            from pptx import Presentation
            from pptx.util import Inches
        except Exception:
            message_critical(self,'PPTX export','python-pptx is unavailable. Run the NIRUPRES installer again to install release dependencies.'); return
        p,_=QFileDialog.getSaveFileName(self,'Export pixel-perfect PPTX',self._dialog_start('presentation.pptx'),'PowerPoint (*.pptx)')
        if not p:return
        self._remember_dialog(p)
        p=p if p.lower().endswith('.pptx') else p+'.pptx'; prs=Presentation(); ar={'16:9':16/9,'16:10':16/10,'4:3':4/3,'A4':297/210}.get(self.data.get('aspect','16:9'),16/9); prs.slide_width=Inches(13.333); prs.slide_height=int(prs.slide_width/ar)
        with tempfile.TemporaryDirectory(prefix='nirupres-pptx-') as td:
            for i,s in enumerate(self.data['slides']):
                if s.get('hidden'): continue
                img=self._render_export_image(s,i,*EXPORT_PROFILES['pptx']); fn=str(Path(td)/f'{i:03d}.png'); img.save(fn)
                slide=prs.slides.add_slide(prs.slide_layouts[6]); slide.shapes.add_picture(fn,0,0,width=prs.slide_width,height=prs.slide_height)
                notes=s.get('notes','').strip()
                if notes:
                    try: slide.notes_slide.notes_text_frame.text=notes
                    except Exception: pass
            prs.save(p)
        self.statusBar().showMessage(f'Pixel-perfect PPTX exported: {p}',3500)
    def asset_manager(self):
        usage={}
        for i,s in enumerate(self.data.get('slides',[])):
            for field,label in (('image','IMAGE'),('backgroundImage','BACKGROUND'),('videoPoster','VIDEO POSTER'),('videoSource','VIDEO')):
                path=s.get(field,'')
                if path: usage.setdefault(path,{'slides':[],'roles':set()}); usage[path]['slides'].append(i+1); usage[path]['roles'].add(label)
        total_bytes=sum(Path(path).stat().st_size for path in usage if Path(path).exists())
        d=QDialog(self); d.setWindowTitle('Assets'); d.resize(760,500); v=QVBoxLayout(d); v.addWidget(QLabel(f'{len(usage)} unique image asset(s) · {total_bytes/MB:.1f} MB current source data · embed mode {self.file_prefs.get("assetQuality","OPTIMIZED")}'))
        lst=QListWidget(); v.addWidget(lst,1)
        for path,meta in usage.items():
            ok=Path(path).exists(); dims=_image_dimensions(path) if ok else (0,0); size=(Path(path).stat().st_size/MB) if ok else 0; roles='/'.join(sorted(meta['roles'])); slides=meta['slides']; it=QListWidgetItem(('✓ ' if ok else '! MISSING  ')+Path(path).name+f'    · {roles} · {dims[0]}×{dims[1]} · {size:.1f} MB · slides {", ".join(map(str,slides))}'); it.setData(Qt.UserRole,path); lst.addItem(it)
        row=QHBoxLayout(); rep=QPushButton('REPLACE SELECTED…'); row.addWidget(rep); row.addStretch(); close=QPushButton('CLOSE'); row.addWidget(close); v.addLayout(row); close.clicked.connect(d.accept)
        def replace():
            it=lst.currentItem()
            if not it:return
            old=it.data(Qt.UserRole); new,_=QFileDialog.getOpenFileName(d,'Replace image',self._dialog_start(),'Images (*.png *.jpg *.jpeg *.webp *.bmp)')
            if not new:return
            self._remember_dialog(new)
            for sl in self.data['slides']:
                if sl.get('image')==old: sl['image']=new
                if sl.get('backgroundImage')==old: sl['backgroundImage']=new
                if sl.get('videoPoster')==old: sl['videoPoster']=new
                if sl.get('videoSource')==old: sl['videoSource']=new
            self.mark_dirty(); self.refresh(self.list.currentRow()); d.accept(); self.asset_manager()
        rep.clicked.connect(replace); d.exec()
    def export_pdf(self):
        p,_=QFileDialog.getSaveFileName(self,'Export PDF',self._dialog_start('presentation.pdf'),'PDF (*.pdf)')
        if not p:return
        self._remember_dialog(p)
        p=p if p.endswith('.pdf') else p+'.pdf'; writer=QPdfWriter(p);writer.setPageSize(QPageSize(QPageSize.A4));writer.setPageOrientation(QPageLayout.Landscape);writer.setResolution(144); painter=QPainter(writer)
        try:
            for i,s in enumerate(self.data['slides']):
                if i:writer.newPage()
                page=QRectF(0,0,writer.width(),writer.height()); ar={'16:9':16/9,'16:10':16/10,'4:3':4/3,'A4':297/210}.get(self.data.get('aspect','16:9'),16/9)
                if page.width()/page.height()>ar:nw=page.height()*ar;r=QRectF((page.width()-nw)/2,0,nw,page.height())
                else:nh=page.width()/ar;r=QRectF(0,(page.height()-nh)/2,page.width(),nh)
                painter.fillRect(page,QColor(THEMES[self.data['theme']]['bg']));render_slide(painter,r,s,self.data['theme'],i,len(self.data['slides']),self.data.get('showNumbers',False),self.data.get('fontFamily','Noto Sans'),self.data.get('showLogo',True),self.data.get('footerText',''),self.data.get('footerAlign','LEFT'),self.data.get('footerSize','MEDIUM'))
        finally:painter.end()
        self.statusBar().showMessage(f'PDF exported: {p}',3000)
    def autosave(self):
        if not self.dirty:return
        try:
            payload=json.dumps({'source':str(self.path) if self.path else '', 'data':self.data},ensure_ascii=False)
            tmp=AUTOSAVE.with_suffix('.tmp'); tmp.write_text(payload,encoding='utf-8'); tmp.replace(AUTOSAVE); self.state.setText('● AUTOSAVED')
        except Exception as e:
            self.statusBar().showMessage(f'Autosave failed: {e}',4000)
    def clear_autosave(self):
        try:AUTOSAVE.unlink(missing_ok=True)
        except:pass
    def offer_recovery(self):
        if not AUTOSAVE.exists():return
        try:
            p=json.loads(AUTOSAVE.read_text(encoding='utf-8'))
            if message_question(self,'Recover autosave','Unsaved NIRUPRES work was found. Recover it?',QMessageBox.Yes|QMessageBox.No,QMessageBox.Yes)==QMessageBox.Yes:
                self.data=p['data'];normalize_themes(self.data);self.data.setdefault('theme','B&W');self.data.setdefault('showNumbers',False); self.data.setdefault('showLogo',True); self.data.setdefault('footerText',''); self.data.setdefault('footerAlign','LEFT'); self.data.setdefault('footerSize','MEDIUM'); self.data['footerAlign']=normalize_footer_align(self.data.get('footerAlign')); self.data['footerSize']=normalize_footer_size(self.data.get('footerSize')); self.data.setdefault('fontFamily','Noto Sans'); self.data.setdefault('aspect','16:9'); self.data.setdefault('transition','NONE'); self.data.setdefault('reduceMotion',False); self.data.setdefault('cinematicOpen',True); self.data.setdefault('cinematicClose',True); self.data['transition']=normalize_transition(self.data.get('transition','NONE'));self.data['slides']=[slide_defaults(x) for x in self.data.get('slides',[])];self.path=Path(p['source']) if p.get('source') else None;self.dirty=True;self.refresh();self.history=[];self.history_pos=-1;self.push_history();self.state.setText('● RECOVERED')
            else:self.clear_autosave()
        except:self.clear_autosave()
    def confirm_discard(self):
        if not self.dirty:return True
        r=message_question(self,'Unsaved changes','Save changes before continuing?',QMessageBox.Save|QMessageBox.Discard|QMessageBox.Cancel,QMessageBox.Save)
        if r==QMessageBox.Save:return bool(self.save())
        return r==QMessageBox.Discard
    def closeEvent(self,e):
        if self.confirm_discard():
            if self.presenter is not None:
                try:self.presenter.close_session()
                except Exception:pass
            self.clear_autosave();self.extract.cleanup();e.accept()
        else:e.ignore()
    def toggle_outline(self,on):
        self.outline_mode=bool(on); self.refresh(max(0,self.list.currentRow()))
    def import_markdown(self):
        p,_=QFileDialog.getOpenFileName(self,'Import NIRUPRES Markdown',self._dialog_start(prefer_current=False),'Markdown (*.md *.markdown)')
        if not p:return
        self._remember_dialog(p)
        if self.dirty and not self.confirm_discard():return
        try:
            self.data,report=import_nirupres_markdown_report(Path(p).read_text(encoding='utf-8'),Path(p).parent); self.path=None; self.dirty=True; self.refresh(); self.history=[];self.history_pos=-1;self.push_history(); self.statusBar().showMessage(f'Imported {len(self.data["slides"])} slides · {len(report)} normalization(s)',4000); self.validate_deck(silent=True); report and message_info(self,'Markdown import',f'Imported {len(self.data["slides"])} slides.\n\nNormalized:\n'+'\n'.join('• '+x for x in report[:12]))
        except Exception as e: message_critical(self,'Markdown import failed',str(e))
    def export_markdown(self):
        p,_=QFileDialog.getSaveFileName(self,'Export NIRUPRES Markdown',self._dialog_start('presentation.md'),'Markdown (*.md)')
        if not p:return
        self._remember_dialog(p)
        p=p if p.lower().endswith('.md') else p+'.md'; Path(p).write_text(export_nirupres_markdown(self.data),encoding='utf-8'); self.statusBar().showMessage(f'Markdown exported: {p}',3000)
    def export_ai_template(self):
        p,_=QFileDialog.getSaveFileName(self,'Export AI template',self._dialog_start('NIRUPRES-AI-TEMPLATE.md'),'Markdown (*.md)')
        if not p:return
        self._remember_dialog(p)
        p=p if p.lower().endswith('.md') else p+'.md'; Path(p).write_text(AI_TEMPLATE,encoding='utf-8'); self.statusBar().showMessage(f'AI template exported: {p}',3000)
    def export_ai_brief(self):
        p,_=QFileDialog.getSaveFileName(self,'Export AI brief',self._dialog_start('NIRUPRES-AI-BRIEF.md'),'Markdown (*.md)')
        if not p:return
        self._remember_dialog(p)
        p=p if p.lower().endswith('.md') else p+'.md'; Path(p).write_text(export_ai_brief(self.data),encoding='utf-8'); self.statusBar().showMessage(f'AI brief exported: {p}',3000)
    def copy_ai_prompt(self):
        QApplication.clipboard().setText(export_ai_brief(self.data)); self.statusBar().showMessage('AI prompt + current deck copied to clipboard',3000)
    def validate_deck(self,silent=False):
        issues=[]
        for i,s in enumerate(self.data.get('slides',[])):
            n=i+1; words=len((s.get('title','')+' '+s.get('body','')).split())
            if s.get('layout') not in LAYOUTS: issues.append((i,'ERROR',f'Unknown layout: {s.get("layout")}'))
            if not (s.get('title') or s.get('body') or s.get('image') or s.get('imageRequest')): issues.append((i,'WARN','Empty slide'))
            if s.get('image') and not Path(s.get('image')).exists(): issues.append((i,'ERROR',f'Missing image: {s.get("image")}'))
            if s.get('imageRequest') and not s.get('image'): issues.append((i,'TODO',f'Unresolved image request: {s.get("imageRequest")}'))
            if s.get('layout') in ('IMAGE','PHOTO','HERO IMAGE','IMAGE + QUOTE','FULL BLEED') and not (s.get('image') or s.get('imageRequest')): issues.append((i,'WARN',f'{s.get("layout")} layout has no image'))
            if s.get('layout')=='FULL BLEED' and s.get('image') and int(s.get('overlay',0))<12 and (s.get('title') or s.get('body')): issues.append((i,'WARN','FULL BLEED text may need more overlay for image contrast'))
            limit=70 if s.get('layout') in ('TITLE','STATEMENT','PHOTO','HERO IMAGE','IMAGE + QUOTE','BIG NUMBER','DATA / KPI') else 120
            if words>limit: issues.append((i,'WARN',f'Text-heavy slide: {words} words'))
            if len(s.get('title',''))>120: issues.append((i,'WARN','Very long title; may reach minimum text size'))
            eff=normalize_theme(s.get('themeOverride','DECK') if s.get('themeOverride','DECK')!='DECK' else self.data.get('theme','B&W'))
            if THEMES.get(eff,{}).get('uddevalla') and not self.data.get('showLogo',True) and s.get('logoMode','AUTO')=='ON': issues.append((i,'WARN','Slide requests logo ON but deck LOGO master switch is OFF'))
            if s.get('accentOverride','AUTO')!='AUTO' and not THEMES.get(eff,{}).get('uddevalla'): issues.append((i,'WARN','Uddevalla accent override is set on a non-Uddevalla theme'))
            if s.get('layout')=='FULL BLEED' and s.get('logoMode','AUTO')=='ON' and not s.get('image'): issues.append((i,'WARN','FULL BLEED logo is ON but no background image is set'))
        if silent:
            if issues:self.statusBar().showMessage(f'{len(issues)} deck health note(s) · Ctrl+K → Deck health',4500)
            return issues
        d=QDialog(self); d.setWindowTitle('Deck Health'); d.resize(720,520); v=QVBoxLayout(d)
        errors=sum(1 for x in issues if x[1]=='ERROR'); todos=sum(1 for x in issues if x[1]=='TODO'); warns=sum(1 for x in issues if x[1]=='WARN')
        v.addWidget(QLabel(f'{len(self.data.get("slides",[]))} slides  ·  {errors} errors  ·  {warns} warnings  ·  {todos} image requests'))
        if not issues: v.addWidget(QLabel('✓ No issues found.'))
        else:
            for i,kind,msg in issues:
                b=QPushButton(f'{kind}   {i+1:02d}   {msg}'); b.setToolTip('Jump to slide'); b.clicked.connect(lambda _=False,n=i:(self.list.setCurrentRow(n),d.accept())); v.addWidget(b)
            v.addStretch()
        bb=neutralize_button_box(QDialogButtonBox(QDialogButtonBox.Close));bb.rejected.connect(d.reject);v.addWidget(bb);d.exec(); return issues
    def presenter_view(self):
        screens=QApplication.screens()
        if len(screens)<2:
            message_info(self,'Presenter View','Only one display is currently available. Starting normal presentation instead.'); self.present(False); return
        d=DisplaySetup(self)
        if d.exec()!=QDialog.Accepted:return
        ps,aus=d.selection()
        try:
            DISPLAY_PREFS.write_text(json.dumps({'presenter':ps.name(),'audience':aus.name()}),encoding='utf-8')
        except:pass
        self.present(False,ps,aus)
    def online_present(self,from_start=False):
        # Windowed Audience output for Teams/Meet/Zoom. No service integration is
        # required: share the window titled "NIRUPRES — Audience · SHARE THIS WINDOW".
        screen=self.screen() or QApplication.primaryScreen()
        self.present(from_start,screen,screen,online=True)
        if self.presenter is not None:
            self.statusBar().showMessage('ONLINE: share only the NIRUPRES — Audience window',5000)
    def presentation_overview(self):
        d=QDialog(self);d.setWindowTitle('Presentation overview');d.resize(980,650);v=QVBoxLayout(d);v.addWidget(QLabel(f'{self.data.get("title","Untitled")}  ·  {len(self.data.get("slides",[]))} slides'))
        grid=QGridLayout();v.addLayout(grid,1)
        for i,s in enumerate(self.data.get('slides',[])):
            box=QPushButton(f'{i+1:02d}  {s.get("title") or "(untitled)"}\n{s.get("layout","TITLE")}'+('  · HIDDEN' if s.get('hidden') else ''));box.setMinimumSize(180,80);box.clicked.connect(lambda _=False,n=i:(self.list.setCurrentRow(n),d.accept()));grid.addWidget(box,i//4,i%4)
        bb=neutralize_button_box(QDialogButtonBox(QDialogButtonBox.Close));bb.rejected.connect(d.reject);v.addWidget(bb);d.exec()
    def _presentation_finished(self):
        self.presenter=None; self.show(); self.raise_(); self.activateWindow(); self.setFocus(Qt.OtherFocusReason)
    def present(self,from_start=False,presenter_screen=None,audience_screen=None,online=False):
        if self.presenter is not None:
            self.presenter.close_session(); self.presenter=None
        slides=self.data.get('slides',[])
        if not slides:return
        if not any(not s.get('hidden',False) for s in slides):
            message_warning(self,'Cannot present','All slides are hidden. Unhide at least one slide first.'); return
        start=0 if from_start else max(0,self.list.currentRow())
        screens=QApplication.screens(); primary=QApplication.primaryScreen()
        if presenter_screen is None: presenter_screen=self.screen() or primary
        if audience_screen is None:
            # Reuse a remembered audience display when available; otherwise choose another connected screen.
            prefs={}
            try:prefs=json.loads(DISPLAY_PREFS.read_text(encoding='utf-8')) if DISPLAY_PREFS.exists() else {}
            except:pass
            audience_screen=next((x for x in screens if x.name()==prefs.get('audience')),None) or next((x for x in screens if x!=presenter_screen),presenter_screen)
        self.presenter=Presenter(self.data,start,presenter_screen,audience_screen,online=online); self.presenter.finished.connect(self._presentation_finished)

ASSET_DIR=Path(__file__).resolve().parent/'assets'
CHEVRON=(ASSET_DIR/'chevron-down.svg').as_posix()
CHECKMARK=(ASSET_DIR/'check.svg').as_posix()

STYLE='''
QWidget{background:#080808;color:#e8e8e8;font-family:"Noto Sans",sans-serif;font-size:13px} QMainWindow{background:#080808}
QToolBar{background:#0b0b0b;border:0;border-bottom:1px solid #242424;padding:7px;spacing:4px} QToolButton,QPushButton{qproperty-iconSize:0px 0px;background:#121212;border:1px solid #303030;padding:7px 11px;font-weight:650;min-height:18px} QToolButton:hover,QPushButton:hover{background:#202020;border-color:#555} QPushButton#present{background:#e8e8e8;color:#090909;border-color:#e8e8e8;font-weight:900;padding-right:14px} QPushButton#presentAux{background:#151515;color:#e8e8e8;border:1px solid #3a3a3a;font-weight:850;padding:7px 5px} QPushButton#presentAux:hover{background:#242424;border-color:#777}
QListWidget,QLineEdit,QTextEdit,QComboBox{background:#0e0e0e;border:1px solid #292929;padding:8px;selection-background-color:#eee;selection-color:#090909} QComboBox{padding-right:34px;min-height:18px} QListWidget{padding:5px} QListWidget::item{border:1px solid transparent} QListWidget::item:selected{border:1px solid #777;background:#151515}
QLabel{color:#a8a8a8}
QGroupBox{border:0;border-top:1px solid #242424;margin-top:12px;padding:13px 6px 7px 6px;font-weight:750;color:#c6c6c6}
QGroupBox::title{subcontrol-origin:margin;left:2px;padding:0 7px 0 2px;background:#080808}
QGroupBox::indicator{width:0px;height:0px}
QCheckBox{spacing:9px;color:#d2d2d2;min-height:22px}
QCheckBox::indicator{width:14px;height:14px;border:1px solid #737373;background:#111;border-radius:2px}
QCheckBox::indicator:hover{border:1px solid #d7d7d7;background:#191919}
QCheckBox::indicator:focus{border:1px solid #f0f0f0}
QCheckBox::indicator:checked{background:#e8e8e8;border:1px solid #f4f4f4;image:url(__CHECKMARK__)}
QCheckBox::indicator:checked:hover{background:#fff;border-color:#fff}
QCheckBox::indicator:disabled{background:#0b0b0b;border-color:#303030}
QCheckBox:disabled{color:#555}
QLineEdit:focus,QTextEdit:focus,QComboBox:focus,QListWidget:focus{border:1px solid #8c8c8c;background:#111}
QLineEdit:hover,QTextEdit:hover,QComboBox:hover{border-color:#484848}
QPushButton:focus,QToolButton:focus{border:1px solid #9a9a9a}
QComboBox::drop-down{border:0;border-left:1px solid #303030;width:30px;background:#141414} QComboBox::drop-down:hover{background:#222} QComboBox::down-arrow{image:url(__CHEVRON__);width:12px;height:8px} QComboBox QAbstractItemView{background:#111;color:#e8e8e8;border:1px solid #454545;selection-background-color:#e8e8e8;selection-color:#090909;outline:0;padding:4px} QComboBox:disabled{color:#5d5d5d;background:#0b0b0b;border-color:#202020} QComboBox::drop-down:disabled{background:#0b0b0b;border-left-color:#202020} QPushButton#menuButton{padding-right:34px} QPushButton#menuButton::menu-indicator{image:url(__CHEVRON__);subcontrol-origin:padding;subcontrol-position:center right;width:12px;height:8px;right:10px} QMenu{background:#111;color:#e8e8e8;border:1px solid #454545;padding:5px} QMenu::item{padding:7px 28px 7px 10px} QMenu::item:selected{background:#e8e8e8;color:#090909} QMenu::separator{height:1px;background:#303030;margin:5px 8px}
QSlider::groove:horizontal{height:4px;background:#333;border-radius:2px}
QSlider::groove:horizontal:hover{background:#414141}
QSlider::handle:horizontal{width:14px;margin:-5px 0;background:#ddd;border:1px solid #eee;border-radius:7px}
QSlider::handle:horizontal:hover{background:#fff;border-color:#fff}
QSlider:focus QSlider::handle:horizontal{border:1px solid #fff}
#thumbtitle{font-size:10px;color:#aaa}
#doc{font-size:12px;font-weight:700;letter-spacing:.5px;color:#ddd;background:#0b0b0b;border:1px solid #242424;padding:6px 14px;border-radius:3px} #doc:hover{color:#fff;border-color:#505050;background:#111}
#aboutTitle{font-size:28px;font-weight:900;color:#f2f2f2;letter-spacing:2px} #aboutTag{font-size:14px;color:#aaa;padding-bottom:8px}
#deckLabel{font-family:"Noto Sans Mono",monospace;font-size:10px;font-weight:700;color:#737373;letter-spacing:.7px}
#hint{font-family:"Noto Sans Mono",monospace;font-size:10px;color:#666;padding-top:8px}
#presenterLabel{font-family:"Noto Sans Mono",monospace;font-size:10px;font-weight:800;color:#777;letter-spacing:1px;padding:2px 0} #presenterState{font-weight:800;color:#d0d0d0}
QSplitter::handle{background:#181818;width:1px}
QSplitter::handle:hover{background:#4a4a4a}
QMessageBox{background:#080808} QMessageBox QLabel{color:#d8d8d8;min-width:260px} QDialogButtonBox QPushButton{qproperty-iconSize:0px 0px;min-width:112px;min-height:34px}
QToolBar{spacing:4px;border-bottom:1px solid #242424;padding:5px;background:#0d0d0d}
'''
STYLE=STYLE.replace('__CHECKMARK__',CHECKMARK).replace('__CHEVRON__',CHEVRON)
def run():
    app=QApplication(sys.argv);app.setApplicationName(APP);app.setApplicationVersion(VERSION);app.setOrganizationName('Nicklas Rudolfsson'); icon=QIcon(str(ASSETS/'nirupres.svg')); app.setWindowIcon(icon); w=Main(); w.setWindowIcon(icon);w.show();sys.exit(app.exec())
if __name__=='__main__':run()
